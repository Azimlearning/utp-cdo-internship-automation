"""Generate fake-data Excel workbooks for testing the Excel-driven flows, and check them against the flow definitions.

    python tools/testdata.py                          # writes testdata/*.xlsx (no email-sending file)
    python tools/testdata.py --staff-email you@utp.edu.my --semester "Sep 2026"
                                                      # also writes the flow 2 file, whose supervisor emails go to YOU
    python tools/testdata.py --check                  # only verify the files against the definitions in flows/*.zip

Every person, company and address is invented (`@sandbox.test`); nothing here is real student or host-company data.
The generated workbooks are INPUT files (drop files, Excel inputs for the checking flows) and MOCK templates. A mock template is
for sandbox use or as a column reference ONLY: the real templates are cleared in place, never replaced (D-03).

Column names are taken from the flow definitions (trailing spaces and all: 'Course ', 'Host Company '); `--check` re-reads
flows/*.zip and fails if a column a flow reads from its table is missing from the generated header row.
"""
import argparse
import pathlib
import re
import sys
import zipfile
import json

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / "testdata"
SEM = "Sep 2026"
ORG_TYPES = ["government", "government linked company (glc)", "multinational company (mnc)", "private local",
             "statutory body", "university local"]  # a subset of the values flow 2's CountSV step recognises


def write(path, sheets):
    """sheets: {sheet name: [(table name, headers, rows)]} - one table per sheet block, stacked vertically with a gap."""
    wb = Workbook()
    wb.remove(wb.active)
    for sname, tables in sheets.items():
        ws = wb.create_sheet(sname)
        r0 = 1
        for tname, headers, rows in tables:
            ws.append([None] * 0)
            for c, h in enumerate(headers, 1):
                ws.cell(row=r0, column=c, value=h)
            for i, row in enumerate(rows, 1):
                for c, v in enumerate(row, 1):
                    ws.cell(row=r0 + i, column=c, value=v)
            last = r0 + max(len(rows), 1)
            ref = f"A{r0}:{ws.cell(row=last, column=len(headers)).coordinate}"
            t = Table(displayName=tname, ref=ref)
            t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(t)
            for c, h in enumerate(headers, 1):
                ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = max(14, len(str(h)) + 2)
            r0 = last + 3
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ---- contracts (from the definitions) -------------------------------------------------------------------------------
STUDENT_LIST = ["Year", "Course ", "Semester", "Batch", "Intake", "Student ID", "Student Name", "Gender", "Nationality", "Country",
                "Student Program", "CGPA", "Sponsor", "Email Address"]
HCSV_DB = ["HC SV Name", "Host SV Salutation", "HC SV Designation/Position", "HC SV Department/Unit/Section",
           "Student ID", "Student Name", "Email Address"]
HC_DB = ["Host Company ", "Type of Host", "Type of placement", "Organisation Type", "Industry Type", "HC Street Name",
         "HC City / Town", "Postcode", "HC State / Country", "Student ID", "Student Name", "Email Address"]
ACK_DB = ["No", "Year", "Course ", "Semester", "Student ID", "Student Name", "Email Address", "Host Company ",
          "Type of placement", "Organisation Type", "HC SV Name", "HC Supervisor Email"]
HCSV_TPL = ["HC Supervisor Name", "Host SV Salutation", "Host SV Designation/Position", "Host SV Department/Unit/Section", "Status",
            "Remarks", "Student ID", "Student Name", "Email Address"]  # note 'Host SV', not 'HC SV' (as flow 12 writes them)
HC_TPL = ["Host Company ( HC ) Name", "Type of placement", "Type of Host", "Organisation Type", "Industry Type", "HC Street Name",
          "HC City / Town", "Postcode", "HC State / Country", "Status", "Remarks", "Student ID", "Student Name", "Email Address"]

# real-system test inputs for flows 12/13 line up with these SharePoint test items (add them to the list first, delete after):
#   HC Supervisor list:  'ZZ TEST Exact SV' (Mr, Manager, HR)   'ZZ TEST Partial SV' (Ms, Manager, Finance)
#   Host Company list:   'ZZ TEST Exact Co' (all fields as below) 'ZZ TEST Partial Co' (street '2 Test Road')


def student(i, email=None, **kw):
    d = {"Year": 3, "Course ": "Computer Science", "Semester": SEM, "Batch": "B1", "Intake": "Sep 2023", "Student ID": f"99{i:05d}",
         "Student Name": f"Test Student {i:02d}", "Gender": "Female" if i % 2 else "Male", "Nationality": "Malaysian",
         "Country": "Malaysia", "Student Program": "BSc Computer Science", "CGPA": 3.0 + (i % 10) / 10, "Sponsor": "Self",
         "Email Address": email or f"test.student{i:02d}@sandbox.test"}
    d.update(kw)
    return [d[h] for h in STUDENT_LIST]


def build(staff_email=None, semester=SEM):
    global SEM
    SEM = semester
    files = {}

    # Flow 14: StudentList table. Rows 1-6 are new students; 3 repeats row 2's email (file-internal duplicate: created once);
    # 7 has a different case/spacing of row 1's email (the flow lower-cases and trims); 8 has no email (observe what happens)
    rows = [student(1), student(2), student(3, email="test.student02@sandbox.test"), student(4), student(5), student(6),
            student(7, email="  TEST.Student01@sandbox.test "), student(8, email="")]
    files["StudentList_TEST.xlsx"] = {"Sheet1": [("StudentList", STUDENT_LIST, rows)]}
    # a valid workbook whose table is not called StudentList: the flow's Excel read must fail (a realistic admin mistake)
    files["StudentList_wrongtable.xlsx"] = {"Sheet1": [("Sheet1Table", STUDENT_LIST, rows[:2])]}

    # Flow 12: Database table in a file in MasterList_RawData_ToBeProcess
    def sv(name, sal, pos, dept, i):
        return [name, sal, pos, dept, f"98{i:05d}", f"Test Student {i:02d}", f"test.student{i:02d}@sandbox.test"]
    rows = [sv("ZZ TEST Exact SV", "Mr", "Manager", "HR", 1),            # identical to the list item   -> Matched
            sv("ZZ TEST Partial SV", "Ms", "Director", "Finance", 2),     # same name, other position    -> Partially Not Matched
            sv("ZZ TEST New SV", "Mr", "Engineer", "Ops", 3),             # not in the list              -> Exactly Not Matched
            sv("zz test exact sv ", "mr", "manager", "hr", 4)]            # case/space variant of row 1  -> Matched (flow lower-cases)
    files["RawData_HCSupervisor_TEST.xlsx"] = {"Sheet1": [("Database", HCSV_DB, rows)]}

    # Flow 13
    def co(name, street, i, typ="Company"):
        return [name, typ, "Internship", "private local", "IT", street, "Ipoh", "30000", "Perak",
                f"97{i:05d}", f"Test Student {i:02d}", f"test.student{i:02d}@sandbox.test"]
    rows = [co("ZZ TEST Exact Co", "1 Test Road", 1), co("ZZ TEST Partial Co", "99 Other Road", 2),
            co("ZZ TEST New Co", "5 New Road", 3), co("zz test exact co ", "1 Test Road", 4)]
    files["RawData_HostCompany_TEST.xlsx"] = {"Sheet1": [("Database", HC_DB, rows)]}

    # Mock templates: the table layout flows 3/4/12/13 read and write. SANDBOX / COLUMN REFERENCE ONLY (D-03).
    files["Template_HCSupervisor_MOCK.xlsx"] = {
        "HCSupervisor": [("HCSupervisor", HCSV_TPL, [])], "SVExactlyNotMatched": [("SVExactlyNotMatched", HCSV_TPL, [])]}
    files["Template_HostCompany_MOCK.xlsx"] = {
        "HostCompany": [("HostCompany", HC_TPL, [])], "ExactlyNotMatched": [("ExactlyNotMatched", HC_TPL, [])]}

    # Flows 3/4 test input: rows to type/paste into the *Exactly Not Matched* table of the real template (clear them after)
    # The `_rowfail` copies have a 300-character name in row 2: SharePoint rejects it (255-character limit), so the 2nd item
    # fails after the 1st was created (the partial-write case).
    def sv_added(n, name=None):
        return [name or f"ZZ TEST Added SV {n}", "Mr", "Tester", "QA", "Exactly Not Matched", "[]", f"960000{n:02d}",
                f"Test Student {90 + n}", f"test.student{90 + n}@sandbox.test"]

    def co_added(n, name=None):
        return [name or f"ZZ TEST Added Co {n}", "Internship", "Company", "private local", "IT", f"{n} Test Lane", "Ipoh", "30000",
                "Perak", "Exactly Not Matched", "[]", f"961000{n:02d}", f"Test Student {80 + n}", f"test.student{80 + n}@sandbox.test"]
    files["Rows_for_Flow3_SVExactlyNotMatched.xlsx"] = {"Sheet1": [("Rows", HCSV_TPL, [sv_added(1), sv_added(2)])]}
    files["Rows_for_Flow3_rowfail.xlsx"] = {"Sheet1": [("Rows", HCSV_TPL, [sv_added(1), sv_added(2, "X" * 300)])]}
    files["Rows_for_Flow4_ExactlyNotMatched.xlsx"] = {"Sheet1": [("Rows", HC_TPL, [co_added(1), co_added(2)])]}
    files["Rows_for_Flow4_rowfail.xlsx"] = {"Sheet1": [("Rows", HC_TPL, [co_added(1), co_added(2, "X" * 300)])]}

    # Flow 2: only when the supervisor emails can be pointed at the tester (the flow emails the supervisor)
    if staff_email:
        def ack(n, sem=None, em=None, supem=None, host="ZZ TEST Co A", sv="ZZ TEST SV A", org="private local", empty=None):
            r = [n, 3, "Computer Science", sem or semester, f"95{n:05d}", f"Test Student {n:02d}", em or f"test.student{n:02d}@sandbox.test",
                 host, "Internship", org, sv, supem or staff_email]
            if empty is not None:
                r[empty] = None
            return r
        rows = [ack(1), ack(2, host="ZZ TEST Co B", sv="ZZ TEST SV B", org="government"),
                ack(3, sem="Jan 2020"),                       # wrong semester          -> InsufficientData (CurrentSemesterCheck)
                ack(4, supem="not-an-email"),                  # bad supervisor email    -> InsufficientData (CheckEmailFormat)
                ack(5, empty=7),                               # empty Host Company cell -> InsufficientData (FindEmptyCell)
                ack(6, host="ZZ TEST Co A", sv="ZZ TEST SV A"),  # same host+SV as row 1   -> grouped (same-name rearrange)
                ack(7), ack(7)]                                # exact duplicate pair    -> RemoveDuplicates
        files["RawData_Acknowledgement_TEST.xlsx"] = {"Database": [("Database", ACK_DB, rows)]}
    return files


# ---- check against the definitions ----------------------------------------------------------------------------------
def definition(zip_name):
    z = zipfile.ZipFile(ROOT / "flows" / zip_name)
    n = [x for x in z.namelist() if x.endswith("definition.json")][0]
    return z.read(n).decode("utf-8-sig")


def referenced_columns(zip_name, loop):
    """Columns a flow reads from rows of the table loop `loop`: items('loop')?['col'] and items('loop')?[concat('a','/','b')]."""
    txt = definition(zip_name)
    cols = set(re.findall(r"items\('%s'\)\??\[\\?'([^'\\]+)\\?'\]" % loop, txt))
    for m in re.findall(r"items\('%s'\)\??\[concat\(([^\]]*?)\)\]" % loop, txt):
        cols.add("".join(re.findall(r"\\?'([^'\\]*)\\?'", m)))
    return cols


def headers_of(path, table):
    wb = load_workbook(path)
    for ws in wb.worksheets:
        if table in ws.tables:
            ref = ws.tables[table].ref
            first = ws[ref.split(":")[0]].row
            return [c.value for c in ws[first][:len(ws[ref.split(":")[0]:ref.split(":")[1]][0])]], ws.tables[table]
    return None, None


CHECKS = [  # file, table, flow export, loop variable that iterates the table's rows, columns we deliberately don't supply
    ("StudentList_TEST.xlsx", "StudentList", "[LIVE-CDO] UTPCSDCDO Store Student List into SharePoint List.zip", "Apply_to_each_3", {"EmailAddress"}),
    ("RawData_HCSupervisor_TEST.xlsx", "Database", "[LIVE-CDO] UTPCSDCDO HC Supervisor Info MasterList Checking.zip", "Apply_to_each_2", set()),
    ("RawData_HostCompany_TEST.xlsx", "Database", "[LIVE-CDO] UTPCSDCDO Host Company Info MasterList Checking.zip", "Apply_to_each_2", set()),
    ("Rows_for_Flow3_SVExactlyNotMatched.xlsx", "Rows", "[LIVE-CDO] UTPCSDCDO Add HC Supervisor into MasterList.zip", "Apply_to_each", set()),
    ("Rows_for_Flow4_ExactlyNotMatched.xlsx", "Rows", "[LIVE-CDO] UTPCSDCDO Add Host Company into MasterList.zip", "Apply_to_each", set()),
]


# (mock template, table, flow export, a table id that flow writes to): the headers must equal the flow's add-row columns
TEMPLATE_CHECKS = [
    ("Template_HCSupervisor_MOCK.xlsx", "HCSupervisor", "[LIVE-CDO] UTPCSDCDO HC Supervisor Info MasterList Checking.zip", "05D719C3"),
    ("Template_HCSupervisor_MOCK.xlsx", "SVExactlyNotMatched", "[LIVE-CDO] UTPCSDCDO HC Supervisor Info MasterList Checking.zip", "CB10AEA9"),
    ("Template_HostCompany_MOCK.xlsx", "HostCompany", "[LIVE-CDO] UTPCSDCDO Host Company Info MasterList Checking.zip", "6A8E8DA9"),
    ("Template_HostCompany_MOCK.xlsx", "ExactlyNotMatched", "[LIVE-CDO] UTPCSDCDO Host Company Info MasterList Checking.zip", "08F6F0A2"),
]


def addrow_columns(zip_name, table_id):
    txt = definition(zip_name)
    cols = set()
    for m in re.finditer(r'"table"\s*:\s*"\{%s[^"]*\}"((?:(?!"table").)*?)"item/Email Address"' % table_id, txt, re.S):
        cols |= {c.replace("~1", "/") for c in re.findall(r'"item/([^"]+)"', m.group(0))}
    return cols


def check():
    bad = 0
    for fname, table, zname, tid in TEMPLATE_CHECKS:
        path = OUT / fname
        hdr, _ = headers_of(path, table) if path.exists() else (None, None)
        want = addrow_columns(zname, tid)
        if hdr is None or not want:
            print(f"CHECK: {fname}/{table}: table or flow columns not found"); bad += 1; continue
        diff = sorted(set(want) ^ set(hdr))
        print(f"{fname}: table {table}: flow writes {len(want)} columns; differences: {diff or 'none'}")
        bad += bool(diff)
    for fname, table, zname, loop, skip in CHECKS:
        path = OUT / fname
        if not path.exists():
            print(f"CHECK: {fname} not generated yet"); bad += 1; continue
        hdr, _ = headers_of(path, table)
        if hdr is None:
            print(f"CHECK: {fname}: table {table} not found"); bad += 1; continue
        need = referenced_columns(zname, loop) - skip
        missing = sorted(need - set(hdr))
        print(f"{fname}: table {table}, {len(hdr)} columns; the flow reads {len(need)}; missing: {missing or 'none'}")
        bad += bool(missing)
    return bad == 0


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("--staff-email")
    ap.add_argument("--semester", default=SEM)
    ap.add_argument("--check", action="store_true")
    a = ap.parse_args()
    if not a.check:
        for name, sheets in build(a.staff_email, a.semester).items():
            write(OUT / name, sheets)
            print("wrote", (OUT / name).relative_to(ROOT).as_posix())
        if not a.staff_email:
            print("(flow 2's file skipped: pass --staff-email so its supervisor emails go to you)")
    sys.exit(0 if check() else 1)


if __name__ == "__main__":
    main()
