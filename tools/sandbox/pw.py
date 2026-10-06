"""Generate a one-shot Playwright script for the sandbox, to run with the browser tool:

    python tools/sandbox/pw.py load                    # once per browser session (and after editing lib.js/commands.js)
    python tools/sandbox/pw.py flows [filter]          # list sandbox flows (id, state)
    python tools/sandbox/pw.py import NN baseline|fixed|fixedprobe|harness # import build/<pkg>/flowNN-<pkg>.zip as new, all connections = yours
    python tools/sandbox/pw.py on|off NN baseline|fixed # switch one copy (turning on first switches its twins off)
    python tools/sandbox/pw.py alloff                   # switch every [SANDBOX-*] flow off and verify
    python tools/sandbox/pw.py run NN <scenario> baseline|fixed [--noshot]
                                                        # seed lists, trigger, wait, read the run, switch off, check expectation
    python tools/sandbox/pw.py stage NN <scenario>      # upload that scenario's test workbooks (before importing a real-Excel package)
    python tools/sandbox/pw.py filetest NN <name>       # file-trigger scenario on a harness (FILETESTS in the fix module)
    python tools/sandbox/pw.py report NN                # markdown table of the saved results for flow NN

Each command writes tmp/pw/<cmd>.js and prints:
  1. the script path  -> browser_run_code_unsafe(filename=<path>)
  2. for `run`, the results path -> browser_evaluate(function="() => window.__pwResult", filename=<path>)
`load` installs lib.js + commands.js on the page object once (the browser tool echoes every script it runs, so the big
library is sent once); every other command is a few hundred bytes and refuses to run against a stale library.

Scenarios live in the flow's fix module (tools/fixes/fNN_*.py) as SCENARIOS = {name: {...}}:
  clear: [list titles to empty first]          seed: {list title: [item dicts]}
  trigger: Title of the trigger item ('run')   count: [action names to count per loop iteration]
  countLists: [lists to count afterwards]      expect: {baseline|fixed: {status, failedAction, counts, after}}
Safety: scripts only act on flows named [SANDBOX-*] and only call SharePoint on the sandbox OneDrive site (see lib.js).
"""
import base64
import hashlib
import importlib
import json
import pathlib
import re
import sys

ROOT = pathlib.Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "tools"))
HERE = pathlib.Path(__file__).parent
OUT = ROOT / "tmp" / "pw"
FIXES = {p.stem.split("_", 1)[0][1:]: p.stem for p in (ROOT / "tools" / "fixes").glob("f*.py")}
SITE_PROD = "https://utpmy.sharepoint.com/sites/CareerDevelopmentOfficeCDO"


def cfg():
    c = json.loads((ROOT / "tools" / "sandbox.json").read_text(encoding="utf-8"))
    for k in ("env", "account"):
        if k not in (c.get("pa") or {}):
            raise SystemExit(f"tools/sandbox.json needs pa.{k} (see sandbox.example.json)")
    return c


def fix(num):
    if num not in FIXES:
        raise SystemExit(f"no fix module for flow {num}")
    return importlib.import_module(f"fixes.{FIXES[num]}")


def evidence(num):
    return ROOT / "refdocs" / "issues" / "evidence" / f"flow-{num}"


def mcp_path(p):
    # the Playwright MCP server compares paths case-sensitively against a lowercase drive letter
    s = str(p)
    return s[0].lower() + s[1:] if s[1:2] == ":" else s


def _lib():
    return (HERE / "lib.js").read_text(encoding="utf-8") + "\n" + (HERE / "commands.js").read_text(encoding="utf-8")


def _version():
    return hashlib.sha1(_lib().encode("utf-8")).hexdigest()[:10]


def _write(cmd, js):
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / f"{cmd}.js"
    path.write_text(js, encoding="utf-8")
    print(f"script:  {mcp_path(path)}")
    return path


def load():
    src = _lib()
    body = "\n".join(l for l in src.splitlines() if l.strip() and not l.strip().startswith("//"))
    names = sorted(set(re.findall(r"^async function (cmd_\w+)", src, re.M)))
    v = _version()
    _write("load", "async (page) => {\n" + body + f"\npage.__pw = {{v: '{v}', {', '.join(names)}}};\n"
           f"return 'sandbox library {v} loaded';\n}}\n")


def emit(cmd, entry, params):
    c = cfg()
    P = {"env": c["pa"]["env"], "account": c["pa"]["account"], "site": c["replace"][SITE_PROD],
         "triggerList": c["harness"].get("run_trigger_title", "CDO Sandbox - Flow Run Trigger"), "timeout": 300, **params}
    return _write(cmd, "async (page) => {\n"
                  f"if (!page.__pw || page.__pw.v !== '{_version()}') throw new Error('sandbox library missing or stale: run pw.py load first');\n"
                  f"return await page.__pw.{entry}(page, {json.dumps(P, ensure_ascii=False)});\n}}\n")


def _embed(scenario, content=True):
    """Replace {"file": "X.xlsx"} upload entries by their base64 content (read from testdata/); the browser script cannot read files."""
    for u in scenario.get("upload", []) + scenario.get("drop", []):
        if u.get("b64"):  # inline content (e.g. a deliberately broken file)
            u["name"] = u.get("name") or u["file"]
            continue
        src = ROOT / "testdata" / u["file"]
        u["name"] = u.get("name") or u["file"]
        if content:  # only `stage` carries the bytes; runs refer to the files the page already holds
            u["b64"] = base64.b64encode(src.read_bytes()).decode("ascii")
    return scenario


def flow_name(num, pkg):
    if pkg == "harness":
        return f"[SANDBOX-HARNESS] flow {num} trigger swap"
    if pkg not in ("baseline", "fixed", "fixedprobe", "baselinex", "fixedx"):
        raise SystemExit("package must be baseline, fixed, fixedprobe, baselinex, fixedx or harness")
    return f"[SANDBOX-{pkg.upper()}] {fix(num).DISPLAY_NAME}"


def workbook_check(num, r):
    """Verify the rows REAL Excel wrote into the downloaded template, by Status per table, against the scenario's `workbook`."""
    sc = getattr(fix(num), "SCENARIOS", {}).get(r.get("scenario"), {})
    want = sc.get("workbook")
    files = r.get("files") or {}
    if not want or not files:
        return "—"
    import collections
    import io

    from openpyxl import load_workbook
    got = collections.defaultdict(collections.Counter)
    for key, b64 in files.items():
        if str(b64).startswith("HTTP"):
            return f"❌ download {b64}"
        wb = load_workbook(io.BytesIO(base64.b64decode(b64)))
        for ws in wb.worksheets:
            for tname in ws.tables:
                rows = list(ws[ws.tables[tname].ref])
                hdr = [c.value for c in rows[0]]
                if "Status" not in hdr:
                    continue
                si = hdr.index("Status")
                for row in rows[1:]:
                    if row[si].value:
                        got[tname][row[si].value] += 1
    bad = [f"{t_}: {dict(got[t_])} != {s}" for t_, s in want.items() if dict(got[t_]) != s]
    return "✅ " + "; ".join(f"{t_} {dict(got[t_])}" for t_ in want) if not bad else "❌ " + "; ".join(bad)


def report(num):
    rows = []
    for p in sorted((evidence(num) / "results").glob("*.json")):
        r = json.loads(p.read_text(encoding="utf-8"))
        r = json.loads(r) if isinstance(r, str) else r
        # loop actions report an empty error at action level; the iteration's error is kept under counts[...].firstError
        errs = {k: v.get("firstError") for k, v in r.get("counts", {}).items() if v.get("firstError")}
        failed = "; ".join(f"{a['action']}: {(a['error'] or errs.get(a['action'], ''))[:80]}".rstrip(": ")
                           for a in r.get("failed", []) if not a["action"].startswith(("Try_", "Apply_to_each", "Condition"))) or "—"
        counts = "; ".join(f"{k} " + ", ".join(f"{s}={n}" for s, n in v.items() if isinstance(n, int) and n)
                           for k, v in r.get("counts", {}).items()) or "—"
        verdict = {True: "✅", False: "❌ " + "; ".join(r.get("why", [])), None: "—"}[r.get("pass")]
        rows.append(f"| {r['scenario']} | {r['package']} | **{r.get('status')}** | {failed} | {counts} | {verdict} | {workbook_check(num, r)} | "
                    f"`{p.relative_to(evidence(num)).as_posix()}` |")
    print("| Scenario | Package | Run status | Failed actions | Counts | Matches expectation | Excel rows written | Result file |")
    print("|---|---|---|---|---|---|---|---|")
    print("\n".join(rows) or "| (no results saved yet) |||||||")


def main(a):
    sys.stdout.reconfigure(encoding="utf-8")
    if not a:
        raise SystemExit(__doc__)
    cmd = a[0]
    if cmd == "load":
        load()
    elif cmd == "flows":
        emit("flows", "cmd_flows", {"filter": a[1] if len(a) > 1 else ""})
    elif cmd == "alloff":
        emit("alloff", "cmd_alloff", {})
    elif cmd in ("on", "off"):
        num = a[1].zfill(2)
        emit(f"{cmd}_{num}_{a[2]}", "cmd_state", {"flowName": flow_name(num, a[2]), "on": cmd == "on"})
    elif cmd == "import":
        num, pkg = a[1].zfill(2), a[2]
        z = ROOT / "build" / pkg / f"flow{num}-{pkg}.zip"
        if not z.exists():
            raise SystemExit(f"{z} not built: python tools/build.py {num} {pkg}")
        emit(f"import_{num}_{pkg}", "cmd_import", {"zip": mcp_path(z), "flowName": flow_name(num, pkg)})
    elif cmd == "run":
        num, name, pkg = a[1].zfill(2), a[2], a[3]
        sc = getattr(fix(num), "SCENARIOS", {})
        if name not in sc:
            raise SystemExit(f"flow {num} has no scenario '{name}'. Known: {sorted(sc)}")
        ev = evidence(num)
        (ev / "results").mkdir(parents=True, exist_ok=True)
        shot = None if "--noshot" in a else mcp_path(ev / f"{name}-{pkg}.png")
        emit(f"run_{num}_{name}_{pkg}", "cmd_run", {"flowNo": num, "pkg": pkg, "flowName": flow_name(num, pkg),
                                "scenario": _embed({"name": name, **json.loads(json.dumps(sc[name]))}, content=False), "shot": shot})
        print(f"results: {mcp_path(ev / 'results' / f'{name}-{pkg}.json')}")
    elif cmd == "stage":
        # upload a scenario's test workbooks to the sandbox OneDrive (needed BEFORE importing a real-Excel package)
        num, name = a[1].zfill(2), a[2]
        sc = getattr(fix(num), "SCENARIOS", {})
        if name not in sc:
            raise SystemExit(f"flow {num} has no scenario '{name}'. Known: {sorted(sc)}")
        emit(f"stage_{num}_{name}", "cmd_stage", {"scenario": _embed({"name": name, **json.loads(json.dumps(sc[name]))})})
    elif cmd == "filetest":
        # file-trigger scenarios on a harness: python tools/sandbox/pw.py filetest 02 guard
        num, name = a[1].zfill(2), a[2]
        sc = getattr(fix(num), "FILETESTS", {})
        if name not in sc:
            raise SystemExit(f"flow {num} has no file test '{name}'. Known: {sorted(sc)}")
        ev = evidence(num)
        (ev / "results").mkdir(parents=True, exist_ok=True)
        emit(f"filetest_{num}_{name}", "cmd_filetest", {"flowNo": num, "pkg": "harness", "flowName": flow_name(num, "harness"),
                                                      "scenario": {"name": name, **sc[name]}, "timeout": 300})
        print(f"results: {mcp_path(ev / 'results' / f'{name}-harness.json')}")
    elif cmd == "report":
        report(a[1].zfill(2))
    else:
        raise SystemExit(__doc__)


if __name__ == "__main__":
    main(sys.argv[1:])
